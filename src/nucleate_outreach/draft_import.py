"""Import complete, ready-written email drafts from a CSV or XLSX file.

For emails written outside the app (ChatGPT, Claude Chat, by hand). The file
supplies words only. It can NEVER supply a status, approval, reviewer or
signature: those columns are ignored (and reported), and every imported email
enters the normal "Pending review" queue like any other draft.

Subject and body are kept exactly as written (no trimming, no rewrapping).
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

from .contacts import is_valid_email

MAX_FILE_BYTES = 5_000_000
MAX_ROWS = 1000
MAX_SUBJECT = 200
MAX_BODY = 20_000

# accepted spellings of each column (headers are matched case-insensitively)
ALIASES: dict[str, tuple[str, ...]] = {
    "campaign": ("campaign", "campaign_id", "campaign_name"),
    "contact_id": ("contact_id", "contactid", "id"),
    "recipient_email": ("recipient_email", "email", "recipient", "to"),
    "recipient_name": ("recipient_name", "name"),
    "subject": ("subject",),
    "body": ("body", "email_body", "message"),
    "notes": ("personalization_notes", "source_notes", "personalization_source_notes", "notes"),
}
# Columns we deliberately refuse to read.
IGNORED_PREFIXES = ("status", "approv", "signature", "fingerprint", "reviewer", "sent")

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

EXAMPLE_HEADERS = ["campaign", "contact_id", "recipient_email", "recipient_name", "subject", "body",
                   "personalization_notes"]
EXAMPLE_ROWS = [
    ["example-campaign", "C-001", "alex.example@example.org", "Dr. Alex Example",
     "An invitation to the Nucleate Activator Program",
     "Dear Dr. Example,\n\nYour work on protein engineering caught our attention.\n"
     "We would love to tell you about the Activator Program.\n\nBest regards,\nThe Nucleate team",
     "Mentioned their 2025 review article (fictional)"],
    ["example-campaign", "C-002", "sam.sample@example.org", "Prof. Sam Sample",
     "Quantensensorik & Gründung – ein Gespräch?",
     "Liebe Frau Sample,\n\nwir würden uns freuen, mit Ihnen über Ihre Arbeit zu sprechen.\n\nBeste Grüße\nNucleate",
     "Written in German on purpose (fictional)"],
]


@dataclass(frozen=True)
class DraftRow:
    row: int  # row number in the file (header = row 1)
    campaign: str
    recipient_email: str
    subject: str
    body: str
    contact_id: str = ""
    recipient_name: str = ""
    notes: str = ""


@dataclass(frozen=True)
class RowProblem:
    row: int
    message: str


@dataclass
class ImportPlan:
    """What would happen if the file were imported. Nothing is saved yet."""
    ready: list[DraftRow] = field(default_factory=list)
    errors: list[RowProblem] = field(default_factory=list)
    already_exists: list[RowProblem] = field(default_factory=list)
    file_warnings: list[str] = field(default_factory=list)


def _normalise_header(h: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(h or "").strip().lower()).strip("_")


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value if isinstance(value, str) else str(value)


def read_table(filename: str, data: bytes) -> list[dict[str, str]]:
    """Read a CSV or XLSX upload into rows keyed by normalised header. Raises ValueError with a friendly message."""
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f"The file is larger than {MAX_FILE_BYTES // 1_000_000} MB.")
    lower = filename.lower()
    if lower.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("The CSV is not UTF-8. In Excel choose 'CSV UTF-8' when saving.") from exc
        reader = csv.reader(io.StringIO(text, newline=""))
        try:
            raw_rows = [row for row in reader]
        except csv.Error as exc:
            raise ValueError(f"The CSV could not be read: {exc}") from exc
    elif lower.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("Reading .xlsx needs the 'openpyxl' package: pip install openpyxl") from exc
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            raw_rows = [[_cell(c) for c in row] for row in wb.worksheets[0].iter_rows(values_only=True)]
            wb.close()
        except (zipfile.BadZipFile, KeyError, IndexError, OSError, ValueError) as exc:
            raise ValueError("That does not look like a valid .xlsx workbook.") from exc
    else:
        raise ValueError("Please upload a .csv or .xlsx file.")

    # drop completely empty trailing rows (common in spreadsheets)
    while raw_rows and not any(c.strip() for c in raw_rows[-1]):
        raw_rows.pop()
    if not raw_rows:
        raise ValueError("The file is empty.")
    if len(raw_rows) - 1 > MAX_ROWS:
        raise ValueError(f"The file has more than {MAX_ROWS} rows. Split it into smaller files.")
    headers = [_normalise_header(h) for h in raw_rows[0]]
    return [{"__row__": str(n), **{h: _cell(v) for h, v in zip(headers, row) if h}}
            for n, row in enumerate(raw_rows[1:], start=2)] if headers else []


def _column_map(headers: list[str]) -> tuple[dict[str, str], list[str], list[str]]:
    """Map our field names to the file's header names. Returns (mapping, ignored_headers, problems)."""
    mapping: dict[str, str] = {}
    problems: list[str] = []
    for field_name, names in ALIASES.items():
        found = [h for h in headers if h in names]
        if len(found) > 1:
            problems.append(f"Several columns look like '{field_name}': {', '.join(found)}. Keep only one.")
        if found:
            mapping[field_name] = found[0]
    ignored = [h for h in headers if h.startswith(IGNORED_PREFIXES)]
    return mapping, ignored, problems


def plan_import(filename: str, data: bytes, existing, default_campaign: str = "") -> ImportPlan:
    """Validate a file and work out which rows could be imported.

    `existing(campaign, email) -> bool` tells us whether that person already has a draft there.
    Raises ValueError for problems with the whole file (wrong type, missing columns...).
    """
    rows = read_table(filename, data)
    plan = ImportPlan()
    headers = [k for k in (rows[0] if rows else {}) if k != "__row__"]
    if not rows:
        # header only
        raise ValueError("The file has a header row but no emails.")
    mapping, ignored, problems = _column_map(headers)
    if problems:
        raise ValueError(" ".join(problems))
    missing = [f for f in ("recipient_email", "subject", "body") if f not in mapping]
    if "campaign" not in mapping and not default_campaign.strip():
        missing.insert(0, "campaign (or enter a default campaign below)")
    if missing:
        raise ValueError("The file is missing required column(s): " + ", ".join(missing))
    if ignored:
        plan.file_warnings.append(
            "Ignored column(s): " + ", ".join(ignored)
            + ". Status and approvals can never be imported; every email starts as 'Pending review'.")

    seen: dict[tuple[str, str], int] = {}
    for raw in rows:
        n = int(raw["__row__"])
        get = lambda key: raw.get(mapping.get(key, ""), "")  # noqa: E731
        campaign = get("campaign").strip() or default_campaign.strip()
        email = get("recipient_email").strip().lower()
        subject, body = get("subject"), get("body")
        issues = []
        if not campaign:
            issues.append("no campaign")
        elif len(campaign) > 100:
            issues.append("campaign name is longer than 100 characters")
        if not email:
            issues.append("recipient email is empty")
        elif not is_valid_email(email):
            issues.append(f"invalid recipient email {email!r}")
        if not subject.strip():
            issues.append("subject is empty")
        elif "\n" in subject or "\r" in subject:
            issues.append("subject must be a single line")
        elif len(subject) > MAX_SUBJECT:
            issues.append(f"subject is longer than {MAX_SUBJECT} characters")
        elif _CONTROL.search(subject):
            issues.append("subject contains unsupported control characters")
        if not body.strip():
            issues.append("email body is empty")
        elif len(body) > MAX_BODY:
            issues.append(f"body is longer than {MAX_BODY} characters")
        elif _CONTROL.search(body):
            issues.append("body contains unsupported control characters")
        if issues:
            plan.errors.append(RowProblem(n, "; ".join(issues)))
            continue
        key = (campaign, email)
        if key in seen:
            plan.errors.append(RowProblem(n, f"duplicate of row {seen[key]} (same campaign and recipient)"))
            continue
        seen[key] = n
        if existing(campaign, email):
            plan.already_exists.append(RowProblem(n, f"{email} already has an email in campaign '{campaign}'"))
            continue
        plan.ready.append(DraftRow(n, campaign, email, subject, body, get("contact_id").strip(),
                                   get("recipient_name").strip(), get("notes").strip()))
    return plan


# ---------------------------------------------------------------- example files
def example_csv() -> bytes:
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    writer.writerow(EXAMPLE_HEADERS)
    writer.writerows(EXAMPLE_ROWS)
    return out.getvalue().encode("utf-8-sig")  # BOM so Excel opens umlauts correctly


def example_xlsx() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "drafts"
    ws.append(EXAMPLE_HEADERS)
    for row in EXAMPLE_ROWS:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


EXAMPLE_CONTACTS_CSV = (
    "name,institution,email,research_area,outreach_eligible\n"
    "Dr. Alex Example,Example University,alex.example@example.org,protein engineering,yes\n"
    "Prof. Sam Sample,Sample Institute,sam.sample@example.org,quantum sensing,yes\n"
    "Dr. Robin Placeholder,Placeholder College,robin.placeholder@example.org,soft robotics,no\n"
)
